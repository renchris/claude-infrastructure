import hashlib, time, zipfile, os, sys
from openpyxl import Workbook

def mk(path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    ws["B2"] = 42
    wb.save(path)

mk("a.xlsx"); time.sleep(3); mk("b.xlsx")

def sha(p):
    return hashlib.sha256(open(p,"rb").read()).hexdigest()

print("whole-file sha256")
print(" a.xlsx", sha("a.xlsx"), os.path.getsize("a.xlsx"))
print(" b.xlsx", sha("b.xlsx"), os.path.getsize("b.xlsx"))
print(" identical:", sha("a.xlsx")==sha("b.xlsx"))

def parts(p):
    d={}
    with zipfile.ZipFile(p) as z:
        for i in z.infolist():
            d[i.filename]=(hashlib.sha256(z.read(i.filename)).hexdigest(), i.date_time, i.CRC)
    return d

A,B=parts("a.xlsx"),parts("b.xlsx")
print("\nnames equal:", sorted(A)==sorted(B))
print("\nper-part: name | content-sha equal | zip date_time a -> b")
vol=[]
for n in sorted(A):
    eq = A[n][0]==B[n][0]
    print(f"  {n:34s} {str(eq):5s} {A[n][1]} -> {B[n][1]}")
    if not eq: vol.append(n)
print("\nCONTENT-VOLATILE PARTS:", vol)

def rollup(p, skip_docprops):
    z=zipfile.ZipFile(p)
    h=hashlib.sha256()
    for n in sorted(z.namelist()):
        if skip_docprops and n.startswith("docProps/"): continue
        h.update(n.encode()); h.update(z.read(n))
    return h.hexdigest()

print("\nrollup ALL parts (names+content, no zip metadata):")
print("  a:", rollup("a.xlsx",False)); print("  b:", rollup("b.xlsx",False))
print("  equal:", rollup("a.xlsx",False)==rollup("b.xlsx",False))
print("rollup EXCLUDING docProps/*:")
print("  a:", rollup("a.xlsx",True)); print("  b:", rollup("b.xlsx",True))
print("  equal:", rollup("a.xlsx",True)==rollup("b.xlsx",True))

for n in vol:
    for p in ("a.xlsx","b.xlsx"):
        print(f"\n--- {p}:{n} ---")
        print(zipfile.ZipFile(p).read(n).decode("utf-8","replace")[:1200])
